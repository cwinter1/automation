import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import MetaData, Table, create_engine, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.config import settings
from app.models import ColumnDef, Dataset, RawRow
from app.sanitize import IDENTIFIER_RE, validate_identifier

_SLUG_RE = re.compile(r"[^A-Za-z0-9_]+")


def _coerce_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_xlsx(file_bytes: bytes, sheet_name: str | None = None) -> tuple[list[str], list[dict]]:
    """Parse an uploaded .xlsx file into (columns, rows). First row = headers."""
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    if ws is None:
        raise ValueError("Workbook has no sheets")

    row_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(row_iter)
    except StopIteration:
        raise ValueError("Sheet is empty") from None

    if header_row is None or all(cell is None for cell in header_row):
        raise ValueError("Sheet has no header row")

    columns = [
        (str(cell).strip() if cell is not None and str(cell).strip() else f"column_{i + 1}")
        for i, cell in enumerate(header_row)
    ]

    rows: list[dict] = []
    for raw_row in row_iter:
        if raw_row is None or all(cell is None for cell in raw_row):
            continue
        row = {}
        for col_name, cell in zip(columns, raw_row):
            row[col_name] = _coerce_value(cell)
        rows.append(row)

    return columns, rows


def pull_from_db(connection_string: str, table_name: str, max_rows: int | None = None) -> tuple[list[str], list[dict]]:
    """Pull rows from an external DB table into (columns, rows).

    The connection string is used only within this function's local scope —
    it is never persisted to a model, logged, or included in an error message.
    """
    validate_identifier(table_name)
    max_rows = max_rows or settings.MAX_DB_ROWS

    engine = create_engine(connection_string, pool_pre_ping=True)
    try:
        metadata = MetaData()
        try:
            table = Table(table_name, metadata, autoload_with=engine)
        except SQLAlchemyError:
            raise ValueError(f"Could not find or read table {table_name!r}") from None

        try:
            with engine.connect() as conn:
                result = conn.execute(select(table).limit(max_rows))
                rows = [
                    {col: _coerce_value(val) for col, val in row.items()}
                    for row in result.mappings()
                ]
        except SQLAlchemyError:
            raise ValueError(f"Could not query table {table_name!r}") from None

        columns = [c.name for c in table.columns]
        return columns, rows
    finally:
        engine.dispose()


def make_safe_identifier(name: str, existing: set[str]) -> str:
    slug = _SLUG_RE.sub("_", name or "").strip("_").lower()
    if not slug:
        slug = "col"
    if slug[0].isdigit():
        slug = f"_{slug}"
    slug = slug[:50]

    candidate = slug
    suffix = 2
    while candidate in existing:
        candidate = f"{slug}_{suffix}"[:50]
        suffix += 1

    if not IDENTIFIER_RE.match(candidate):
        raise ValueError(f"Could not derive a safe identifier from {name!r}")

    existing.add(candidate)
    return candidate


def replace_active_dataset(db: Session, source_type: str, columns: list[str], rows: list[dict]) -> Dataset:
    """Wipe any existing dataset and its dependents, then insert a fresh one."""
    db.query(Dataset).delete(synchronize_session=False)
    db.flush()

    dataset = Dataset(source_type=source_type, is_active=True)
    db.add(dataset)
    db.flush()

    existing_safe_names: set[str] = set()
    column_defs = []
    for idx, col_name in enumerate(columns):
        safe_name = make_safe_identifier(col_name, existing_safe_names)
        col_def = ColumnDef(
            dataset_id=dataset.id,
            source_name=col_name,
            safe_name=safe_name,
            order_index=idx,
        )
        db.add(col_def)
        column_defs.append(col_def)
    db.flush()

    for row_index, row in enumerate(rows):
        data = {col_def.safe_name: row.get(col_def.source_name) for col_def in column_defs}
        db.add(RawRow(dataset_id=dataset.id, row_index=row_index, data=data))

    db.commit()
    db.refresh(dataset)
    return dataset
